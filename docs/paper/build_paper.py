"""Build the OpenMycelium research manuscript and architecture figures."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image as RLImage,
    KeepTogether,
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    Preformatted,
    Spacer,
    SimpleDocTemplate,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "OPENMYCELIUM_MYCELIUM_MCCL_IEEE_MANUSCRIPT.md"
ASSETS = ROOT / "assets"
ARTIFACTS = ROOT / "artifacts"
OUTPUT = ARTIFACTS / "OpenMycelium_Mycelium_MCCL_IEEE_Manuscript.docx"
PDF_OUTPUT = ARTIFACTS / "OpenMycelium_Mycelium_MCCL_IEEE_Manuscript.pdf"

# compact_reference_guide preset, resolved exactly from the document skill.
PAGE_MARGIN_IN = 1.0
HEADER_FOOTER_IN = 0.492
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = (80, 80, 120, 120)
BODY_FONT = "Calibri"
BODY_SIZE = 11
BODY_AFTER_PT = 6
BODY_LINE = 1.25
H1 = (16, "2E74B5", 18, 10)
H2 = (13, "2E74B5", 14, 7)
H3 = (12, "1F4D78", 10, 5)
INK = "13212B"
BLUE = "0B73E0"
TEAL = "087F6B"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MUTED = "657180"
GREEN = "0A7A5B"
GOLD = "9A6500"
WHITE = "FFFFFF"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(run, name=BODY_FONT, size=BODY_SIZE, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]):
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"table widths must sum to {CONTENT_WIDTH_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell, *CELL_MARGINS_DXA)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def paragraph_bottom_border(paragraph, color="1B2834", size="18", space="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_field(paragraph, instruction: str):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        paragraph._p.append(node)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    relation_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend((color, underline))
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend((r_pr, text_node))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def create_numbering(document: Document) -> tuple[int, int]:
    numbering = document.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    next_abs = max(existing_abs or [0]) + 1
    next_num = max(existing_num or [0]) + 1

    def add_definition(abstract_id: int, num_id: int, fmt: str, text: str):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        suffix = OxmlElement("w:suff")
        suffix.set(qn("w:val"), "tab")
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.extend((tabs, ind, spacing))
        level.extend((start, num_fmt, lvl_text, suffix, p_pr))
        abstract.append(level)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), str(abstract_id))
        num.append(ref)
        numbering.append(num)

    add_definition(next_abs, next_num, "bullet", "•")
    add_definition(next_abs + 1, next_num + 1, "decimal", "%1.")
    return next_num, next_num + 1


def apply_numbering(paragraph, num_id: int):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend((level, num))
    p_pr.append(num_pr)


def setup_document() -> tuple[Document, int, int]:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(PAGE_MARGIN_IN)
    section.right_margin = Inches(PAGE_MARGIN_IN)
    section.bottom_margin = Inches(PAGE_MARGIN_IN)
    section.left_margin = Inches(PAGE_MARGIN_IN)
    section.header_distance = Inches(HEADER_FOOTER_IN)
    section.footer_distance = Inches(HEADER_FOOTER_IN)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal.font.size = Pt(BODY_SIZE)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(BODY_AFTER_PT)
    normal.paragraph_format.line_spacing = BODY_LINE

    for name, tokens in (("Heading 1", H1), ("Heading 2", H2), ("Heading 3", H3)):
        style = styles[name]
        size, color, before, after = tokens
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.0

    title = styles["Title"]
    title.font.name = BODY_FONT
    title._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = rgb(INK)
    title.paragraph_format.space_before = Pt(10)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.line_spacing = 1.0

    subtitle = styles["Subtitle"]
    subtitle.font.name = BODY_FONT
    subtitle.font.size = Pt(11)
    subtitle.font.color.rgb = rgb(MUTED)
    subtitle.paragraph_format.space_after = Pt(4)

    for name in ("Equation", "Code Block", "Figure Caption", "Callout"):
        if name not in styles:
            styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    equation = styles["Equation"]
    equation.font.name = "Cambria Math"
    equation.font.size = Pt(10.5)
    equation.font.color.rgb = rgb(INK)
    equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.paragraph_format.space_before = Pt(5)
    equation.paragraph_format.space_after = Pt(7)
    equation.paragraph_format.keep_together = True
    code = styles["Code Block"]
    code.font.name = "Consolas"
    code.font.size = Pt(8.5)
    code.font.color.rgb = rgb(INK)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.12)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.0
    caption = styles["Figure Caption"]
    caption.font.name = BODY_FONT
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = rgb(MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(9)
    callout = styles["Callout"]
    callout.font.name = BODY_FONT
    callout.font.size = Pt(9.5)
    callout.font.color.rgb = rgb("3D4B57")
    callout.paragraph_format.left_indent = Inches(0.18)
    callout.paragraph_format.right_indent = Inches(0.18)
    callout.paragraph_format.space_before = Pt(6)
    callout.paragraph_format.space_after = Pt(10)
    callout.paragraph_format.line_spacing = 1.15

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("OPENMYCELIUM RESEARCH  |  TECHNICAL PREPRINT v1.0")
    set_run_font(run, size=8, color=MUTED, bold=True)
    paragraph_bottom_border(header, color="D6DEE6", size="6", space="4")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("OpenMycelium  |  22 August 2026  |  ")
    set_run_font(run, size=8, color=MUTED)
    add_field(footer, "PAGE")
    bullet_id, decimal_id = create_numbering(document)
    return document, bullet_id, decimal_id


def add_inline(paragraph, text: str):
    token = re.compile(r"(\*\*.+?\*\*|`.+?`|https?://\S+)")
    for part in token.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Consolas", size=9.5, color="24506E")
        elif part.startswith("http://") or part.startswith("https://"):
            add_hyperlink(paragraph, part.rstrip(".,)"), part.rstrip(".,)"))
            suffix = part[len(part.rstrip(".,)")):]
            if suffix:
                paragraph.add_run(suffix)
        else:
            run = paragraph.add_run(part)
            set_run_font(run)


def add_table(document: Document, rows: list[list[str]]):
    columns = len(rows[0])
    if columns == 2:
        widths = [2700, 6660]
    elif columns == 3:
        widths = [2160, 3600, 3600]
    elif columns == 4:
        widths = [1500, 2500, 2680, 2680]
    else:
        widths = [CONTENT_WIDTH_DXA // columns] * columns
        widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.05
            add_inline(paragraph, value)
            for run in paragraph.runs:
                set_run_font(run, size=8.5, color=INK, bold=(row_index == 0))
            if row_index == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F8FAFC")
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_image(document: Document, path: Path, alt: str):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run()
    shape = run.add_picture(str(path), width=Inches(6.25))
    shape._inline.docPr.set("descr", alt)


def render_markdown(document: Document, bullet_id: int, decimal_id: int):
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    first_title = True
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if not line:
            index += 1
            continue
        if line.startswith("```"):
            language = line[3:].strip()
            body = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                body.append(lines[index])
                index += 1
            paragraph = document.add_paragraph(style="Code Block")
            if language:
                label = paragraph.add_run(language.upper() + "\n")
                set_run_font(label, name="Consolas", size=7.5, color=TEAL, bold=True)
            run = paragraph.add_run("\n".join(body))
            set_run_font(run, name="Consolas", size=8.5, color=INK)
            p_pr = paragraph._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), LIGHT_GRAY)
            p_pr.append(shd)
            index += 1
            continue
        if line.startswith("| "):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                parts = [part.strip() for part in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", part) for part in parts):
                    rows.append(parts)
                index += 1
            add_table(document, rows)
            continue
        if line.startswith("!["):
            match = re.match(r"!\[(.+?)\]\((.+?)\)", line)
            if match:
                add_image(document, ROOT / match.group(2), match.group(1))
            index += 1
            continue
        if line.startswith("$$") and line.endswith("$$"):
            equation_text = line[2:-2].strip().replace("\\quad", "   ").replace("\\text{", "").replace("}", "")
            equation_text = equation_text.replace("\\ldots", "...").replace("\\approx", "≈").replace("\\land", "∧")
            equation_text = equation_text.replace("\\cup", "∪").replace("\\sum", "Σ").replace("\\max", "max").replace("\\min", "min")
            equation_text = equation_text.replace("\\frac", "frac").replace("\\forall", "∀").replace("\\in", "∈")
            equation_text = equation_text.replace("\\alpha", "α").replace("\\beta", "β").replace("\\kappa", "κ").replace("\\tau", "τ")
            equation_text = equation_text.replace("\\Theta", "Θ").replace("\\rho", "ρ").replace("\\lambda", "λ").replace("\\tag", "")
            equation_text = equation_text.replace("\\[", "[").replace("\\]", "]")
            paragraph = document.add_paragraph(style="Equation")
            run = paragraph.add_run(equation_text)
            set_run_font(run, name="Cambria Math", size=10.5, color=INK)
            index += 1
            continue
        if line.startswith("# "):
            if first_title:
                kicker = document.add_paragraph()
                kicker.paragraph_format.space_before = Pt(10)
                kicker.paragraph_format.space_after = Pt(5)
                run = kicker.add_run("IEEE-STYLE ENGINEERING PREPRINT  |  INVENTION DISCLOSURE")
                set_run_font(run, size=8.5, color=TEAL, bold=True)
                paragraph = document.add_paragraph(style="Title")
                add_inline(paragraph, line[2:])
                for run in paragraph.runs:
                    set_run_font(run, size=25, color=INK, bold=True)
                rule = document.add_paragraph()
                rule.paragraph_format.space_after = Pt(10)
                paragraph_bottom_border(rule, color=TEAL, size="20", space="2")
                first_title = False
            index += 1
            continue
        if line.startswith("## "):
            if line.startswith("## Appendix A") or line == "## References":
                document.add_page_break()
            paragraph = document.add_paragraph(style="Heading 1")
            add_inline(paragraph, line[3:])
            index += 1
            continue
        if line.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            add_inline(paragraph, line[4:])
            index += 1
            continue
        if line.startswith("> "):
            paragraph = document.add_paragraph(style="Callout")
            add_inline(paragraph, line[2:])
            p_pr = paragraph._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "EDF7F4")
            p_pr.append(shd)
            border = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "22")
            left.set(qn("w:space"), "8")
            left.set(qn("w:color"), TEAL)
            border.append(left)
            p_pr.append(border)
            index += 1
            continue
        if line.startswith("**Fig. "):
            paragraph = document.add_paragraph(style="Figure Caption")
            add_inline(paragraph, line)
            index += 1
            continue
        if re.match(r"^\d+\. ", line):
            paragraph = document.add_paragraph()
            apply_numbering(paragraph, decimal_id)
            add_inline(paragraph, re.sub(r"^\d+\. ", "", line))
            index += 1
            continue
        if line.startswith("- "):
            paragraph = document.add_paragraph()
            apply_numbering(paragraph, bullet_id)
            add_inline(paragraph, line[2:])
            index += 1
            continue
        paragraph = document.add_paragraph()
        add_inline(paragraph, line.replace("  ", " "))
        if line.startswith("**Technical manuscript") or line.startswith("**OpenMycelium Project") or line.startswith("**22 August"):
            paragraph.paragraph_format.space_after = Pt(2)
            for run in paragraph.runs:
                set_run_font(run, size=10, color=MUTED, bold=True)
        index += 1


def font(size: int, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def center_text(draw, box, text, text_font, fill, spacing=6):
    x1, y1, x2, y2 = box
    bbox = draw.multiline_textbbox((0, 0), text, font=text_font, spacing=spacing, align="center")
    width, height = bbox[2] - bbox[0], bbox[3] - bbox[1]
    draw.multiline_text(((x1 + x2 - width) / 2, (y1 + y2 - height) / 2), text, font=text_font, fill=fill, spacing=spacing, align="center")


def draw_architecture():
    width, height = 1800, 1000
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title_font, section_font, body_font, tiny = font(42, True), font(26, True), font(21), font(18)
    draw.text((70, 45), "OpenMycelium heterogeneous execution architecture", font=title_font, fill="#13212B")
    layers = [
        (110, 160, 1690, 285, "Experience and governance", "Dashboard  |  CLI/API/MCP  |  Workspaces  |  RBAC  |  Audit  |  Model and release registry", "#EAF3FF", "#0B73E0"),
        (110, 320, 1690, 465, "Mycelium control plane", "Evidence graph  →  memory feasibility  →  minimax placement  →  candidate scoring  →  immutable execution contract", "#E8F6F2", "#087F6B"),
        (110, 500, 1690, 655, "Kubernetes execution plane", "Per-vendor indexed Jobs and services  |  global ranks  |  rendezvous  |  storage  |  lifecycle reconciliation", "#F6F2EA", "#9A6500"),
        (110, 690, 1690, 900, "Communication and hardware", "MCCL contract: portable TCP/Gloo today; qualified native and RDMA progression", "#F3F0FA", "#7055B8"),
    ]
    for x1, y1, x2, y2, heading, text, fill, accent in layers:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=20, fill=fill, outline=accent, width=3)
        draw.rectangle((x1, y1, x1 + 13, y2), fill=accent)
        draw.text((x1 + 42, y1 + 22), heading, font=section_font, fill=accent)
        draw.text((x1 + 42, y1 + 68), text, font=body_font, fill="#24313C")
    boxes = [
        (185, 790, 480, 875, "NVIDIA group\nCUDA + NCCL", "#DCF0FF", "#0B73E0"),
        (555, 790, 850, 875, "AMD group\nROCm + RCCL", "#E6F7EE", "#087F6B"),
        (925, 790, 1220, 875, "Intel / other\noneCCL + plugins", "#FFF2D8", "#9A6500"),
        (1295, 790, 1590, 875, "CPU / portable\nGloo + TCP", "#ECE8F7", "#7055B8"),
    ]
    for x1, y1, x2, y2, text, fill, outline in boxes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=14, fill=fill, outline=outline, width=2)
        center_text(draw, (x1, y1, x2, y2), text, tiny, "#13212B")
    for y1, y2 in ((285, 320), (465, 500), (655, 690)):
        draw.line((900, y1 + 4, 900, y2 - 8), fill="#607080", width=5)
        draw.polygon([(890, y2 - 18), (910, y2 - 18), (900, y2 - 4)], fill="#607080")
    draw.text((1270, 925), "State: PostgreSQL + NATS   Telemetry: Prometheus + Grafana", font=tiny, fill="#657180")
    image.save(ASSETS / "openmycelium_architecture.png", dpi=(180, 180))


def draw_flow():
    width, height = 1800, 820
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title_font, node_font, small = font(40, True), font(22, True), font(18)
    draw.text((70, 45), "Mycelium evidence-gated decision lifecycle", font=title_font, fill="#13212B")
    nodes = [
        (80, 170, 330, 300, "Discover", "device, driver,\nruntime, link"),
        (380, 170, 630, 300, "Qualify", "source, freshness,\nadapter, policy"),
        (680, 170, 930, 300, "Place", "memory-constrained\nminimax"),
        (980, 170, 1230, 300, "Score", "parallel and\ntransport candidates"),
        (1280, 170, 1530, 300, "Compile", "rank, image,\nstorage, network"),
    ]
    colors = ["#EAF3FF", "#E8F6F2", "#FFF3DB", "#F3F0FA", "#E9F2F7"]
    outlines = ["#0B73E0", "#087F6B", "#9A6500", "#7055B8", "#24506E"]
    for idx, (x1, y1, x2, y2, heading, detail) in enumerate(nodes):
        draw.rounded_rectangle((x1, y1, x2, y2), radius=18, fill=colors[idx], outline=outlines[idx], width=3)
        center_text(draw, (x1, y1 + 10, x2, y1 + 58), heading, node_font, outlines[idx])
        center_text(draw, (x1, y1 + 55, x2, y2 - 8), detail, small, "#24313C")
        if idx < len(nodes) - 1:
            draw.line((x2 + 8, 235, nodes[idx + 1][0] - 12, 235), fill="#607080", width=5)
            draw.polygon([(nodes[idx + 1][0] - 22, 225), (nodes[idx + 1][0] - 22, 245), (nodes[idx + 1][0] - 8, 235)], fill="#607080")
    draw.rounded_rectangle((495, 420, 1305, 545), radius=20, fill="#EDF7F4", outline="#087F6B", width=4)
    center_text(draw, (495, 420, 1305, 545), "Admission predicate satisfied?\ncapacity AND driver AND adapter AND transport AND policy AND freshness", node_font, "#075D50")
    draw.line((900, 300, 900, 412), fill="#607080", width=5)
    draw.polygon([(890, 400), (910, 400), (900, 414)], fill="#607080")
    draw.rounded_rectangle((120, 635, 760, 750), radius=18, fill="#FCEBEC", outline="#B23A48", width=3)
    center_text(draw, (120, 635, 760, 750), "NO  →  blocked plan\nmissing evidence remains visible and auditable", node_font, "#8C2633")
    draw.rounded_rectangle((1040, 635, 1680, 750), radius=18, fill="#E8F6F2", outline="#087F6B", width=3)
    center_text(draw, (1040, 635, 1680, 750), "YES  →  deploy and reconcile\nportable or qualified direct transport", node_font, "#075D50")
    draw.line((700, 545, 440, 625), fill="#B23A48", width=5)
    draw.polygon([(425, 625), (445, 614), (447, 635)], fill="#B23A48")
    draw.line((1100, 545, 1360, 625), fill="#087F6B", width=5)
    draw.polygon([(1375, 625), (1353, 614), (1353, 635)], fill="#087F6B")
    image.save(ASSETS / "mycelium_decision_flow.png", dpi=(180, 180))


def pdf_font_names() -> tuple[str, str]:
    regular = Path("C:/Windows/Fonts/arial.ttf")
    bold = Path("C:/Windows/Fonts/arialbd.ttf")
    if not regular.exists():
        regular = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
        bold = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")
    if regular.exists() and bold.exists():
        pdfmetrics.registerFont(TTFont("OMSans", str(regular)))
        pdfmetrics.registerFont(TTFont("OMSans-Bold", str(bold)))
        return "OMSans", "OMSans-Bold"
    return "Helvetica", "Helvetica-Bold"


def pdf_markup(text: str) -> str:
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"`(.+?)`", r"<font name='Courier' color='#24506E'>\1</font>", text)
    text = re.sub(r"(https?://[^\s]+)", r"<link href='\1' color='#0B73E0'>\1</link>", text)
    return text


def readable_equation(line: str) -> str:
    value = line[2:-2].strip()
    replacements = {
        "\\quad": "   ", "\\text{": "", "\\ldots": "...", "\\approx": "≈",
        "\\land": "∧", "\\cup": "∪", "\\sum": "Σ", "\\max": "max", "\\min": "min",
        "\\frac": "frac", "\\forall": "∀", "\\in": "∈", "\\alpha": "α", "\\beta": "β",
        "\\kappa": "κ", "\\tau": "τ", "\\Theta": "Θ", "\\rho": "ρ", "\\lambda": "λ",
        "\\tag": "", "\\[": "[", "\\]": "]",
    }
    for source, target in replacements.items():
        value = value.replace(source, target)
    return value.replace("}", "")


def build_pdf():
    regular, bold = pdf_font_names()
    styles = getSampleStyleSheet()
    body = ParagraphStyle("OMBody", parent=styles["BodyText"], fontName=regular, fontSize=9.2, leading=12.2, textColor=colors.HexColor("#13212B"), spaceAfter=6, alignment=TA_JUSTIFY)
    h1 = ParagraphStyle("OMH1", parent=body, fontName=bold, fontSize=14, leading=17, textColor=colors.HexColor("#2E74B5"), spaceBefore=14, spaceAfter=8, keepWithNext=True)
    h2 = ParagraphStyle("OMH2", parent=body, fontName=bold, fontSize=11.5, leading=14, textColor=colors.HexColor("#2E74B5"), spaceBefore=11, spaceAfter=6, keepWithNext=True)
    title = ParagraphStyle("OMTitle", parent=body, fontName=bold, fontSize=22, leading=25, textColor=colors.HexColor("#13212B"), spaceBefore=12, spaceAfter=10, alignment=TA_LEFT)
    kicker = ParagraphStyle("OMKicker", parent=body, fontName=bold, fontSize=8, leading=10, textColor=colors.HexColor("#087F6B"), spaceAfter=7)
    meta = ParagraphStyle("OMMeta", parent=body, fontName=bold, fontSize=8.5, leading=10, textColor=colors.HexColor("#657180"), spaceAfter=5)
    note = ParagraphStyle("OMNote", parent=body, fontSize=8.4, leading=11, leftIndent=12, rightIndent=12, borderColor=colors.HexColor("#087F6B"), borderWidth=1, borderPadding=8, backColor=colors.HexColor("#EDF7F4"), spaceBefore=5, spaceAfter=10)
    equation = ParagraphStyle("OMEquation", parent=body, fontSize=9.2, leading=12, alignment=TA_CENTER, spaceBefore=4, spaceAfter=7)
    caption = ParagraphStyle("OMCaption", parent=body, fontSize=8, leading=10, textColor=colors.HexColor("#657180"), alignment=TA_CENTER, spaceAfter=8)
    code_style = ParagraphStyle("OMCode", parent=body, fontName="Courier", fontSize=7.4, leading=9.2, leftIndent=10, rightIndent=8, borderPadding=7, backColor=colors.HexColor("#F2F4F7"), spaceBefore=3, spaceAfter=8)
    ref_style = ParagraphStyle("OMRef", parent=body, fontSize=7.8, leading=10, leftIndent=14, firstLineIndent=-14, spaceAfter=5, alignment=TA_LEFT)
    cell = ParagraphStyle("OMCell", parent=body, fontSize=7.5, leading=9, spaceAfter=0, alignment=TA_LEFT)
    cell_head = ParagraphStyle("OMCellHead", parent=cell, fontName=bold, textColor=colors.HexColor("#13212B"))

    doc = SimpleDocTemplate(str(PDF_OUTPUT), pagesize=letter, leftMargin=1 * inch, rightMargin=1 * inch, topMargin=0.8 * inch, bottomMargin=0.72 * inch, title="OpenMycelium Mycelium and MCCL Manuscript", author="OpenMycelium Project")

    def page_furniture(canvas, document):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#D6DEE6"))
        canvas.setLineWidth(0.5)
        canvas.line(1 * inch, letter[1] - 0.53 * inch, letter[0] - 1 * inch, letter[1] - 0.53 * inch)
        canvas.setFont(regular, 7.2)
        canvas.setFillColor(colors.HexColor("#657180"))
        canvas.drawString(1 * inch, letter[1] - 0.43 * inch, "OPENMYCELIUM RESEARCH  |  TECHNICAL PREPRINT v1.0")
        canvas.drawRightString(letter[0] - 1 * inch, 0.42 * inch, f"OpenMycelium  |  22 August 2026  |  {document.page}")
        canvas.restoreState()

    story = []
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    first_title = True
    in_references = False
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("```"):
            language = line[3:].strip()
            block = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            if language:
                block.insert(0, language.upper())
            story.append(Preformatted("\n".join(block), code_style))
            index += 1
            continue
        if line.startswith("| "):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                values = [part.strip() for part in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", part) for part in values):
                    rows.append(values)
                index += 1
            columns = len(rows[0])
            col_widths = {2: [1.75 * inch, 4.75 * inch], 3: [1.5 * inch, 2.5 * inch, 2.5 * inch], 4: [1.05 * inch, 1.75 * inch, 1.85 * inch, 1.85 * inch]}.get(columns, [6.5 * inch / columns] * columns)
            data = [[Paragraph(pdf_markup(value), cell_head if r == 0 else cell) for value in row] for r, row in enumerate(rows)]
            table = Table(data, colWidths=col_widths, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B9C4CF")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            story.extend([table, Spacer(1, 6)])
            continue
        if line.startswith("!["):
            match = re.match(r"!\[(.+?)\]\((.+?)\)", line)
            figure = RLImage(str(ROOT / match.group(2)), width=6.35 * inch, height=3.53 * inch) if match else Spacer(1, 1)
            caption_text = ""
            if index + 2 < len(lines) and lines[index + 2].strip().startswith("**Fig. "):
                caption_text = lines[index + 2].strip()
                index += 2
            group = [figure]
            if caption_text:
                group.extend([Spacer(1, 3), Paragraph(pdf_markup(caption_text), caption)])
            story.append(KeepTogether(group))
            index += 1
            continue
        if line.startswith("$$") and line.endswith("$$"):
            story.append(Paragraph(pdf_markup(readable_equation(line)), equation))
            index += 1
            continue
        if line.startswith("# ") and first_title:
            story.append(Paragraph("IEEE-STYLE ENGINEERING PREPRINT  |  INVENTION DISCLOSURE", kicker))
            story.append(Paragraph(pdf_markup(line[2:]), title))
            story.append(Table([[""]], colWidths=[6.5 * inch], rowHeights=[3], style=TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#087F6B"))])))
            story.append(Spacer(1, 10))
            first_title = False
            index += 1
            continue
        if line.startswith("## "):
            in_references = line == "## References"
            if line.startswith("## Appendix A") or in_references:
                story.append(PageBreak())
            story.append(Paragraph(pdf_markup(line[3:]), h1))
            index += 1
            continue
        if line.startswith("### "):
            story.append(Paragraph(pdf_markup(line[4:]), h2))
            index += 1
            continue
        if line.startswith("> "):
            story.append(Paragraph(pdf_markup(line[2:]), note))
            index += 1
            continue
        if re.match(r"^\d+\. ", line):
            items = []
            while index < len(lines) and re.match(r"^\d+\. ", lines[index].strip()):
                item = re.sub(r"^\d+\. ", "", lines[index].strip())
                items.append(ListItem(Paragraph(pdf_markup(item), body), leftIndent=14))
                index += 1
            story.append(ListFlowable(items, bulletType="1", start="1", leftIndent=20, bulletFontName=regular, bulletFontSize=8.5, spaceAfter=5))
            continue
        if line.startswith("- "):
            items = []
            while index < len(lines) and lines[index].strip().startswith("- "):
                items.append(ListItem(Paragraph(pdf_markup(lines[index].strip()[2:]), body), leftIndent=14))
                index += 1
            story.append(ListFlowable(items, bulletType="bullet", leftIndent=20, bulletFontName=regular, bulletFontSize=8, spaceAfter=5))
            continue
        selected_style = ref_style if in_references and line.startswith("[") else body
        if line.startswith("**Technical manuscript") or line.startswith("**OpenMycelium Project") or line.startswith("**22 August"):
            selected_style = meta
        story.append(Paragraph(pdf_markup(line.replace("  ", " ")), selected_style))
        if line.startswith("**22 August"):
            story.append(Spacer(1, 5))
        index += 1
    doc.build(story, onFirstPage=page_furniture, onLaterPages=page_furniture)
    print(PDF_OUTPUT)


def audit_document(document: Document):
    section = document.sections[0]
    assert round(section.left_margin.inches, 3) == PAGE_MARGIN_IN
    assert round(section.right_margin.inches, 3) == PAGE_MARGIN_IN
    assert round(section.top_margin.inches, 3) == PAGE_MARGIN_IN
    assert round(section.bottom_margin.inches, 3) == PAGE_MARGIN_IN
    for table in document.tables:
        grid_widths = [int(col.get(qn("w:w"))) for col in table._tbl.tblGrid]
        assert sum(grid_widths) == CONTENT_WIDTH_DXA
        assert table._tbl.tblPr.find(qn("w:tblInd")).get(qn("w:w")) == str(TABLE_INDENT_DXA)


def main():
    ASSETS.mkdir(parents=True, exist_ok=True)
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    draw_architecture()
    draw_flow()
    document, bullet_id, decimal_id = setup_document()
    render_markdown(document, bullet_id, decimal_id)
    audit_document(document)
    document.core_properties.title = "OpenMycelium: Evidence-Gated Heterogeneous AI Placement and Collective Execution"
    document.core_properties.subject = "Mycelium algorithm, MCCL architecture, market gap, business use cases, and invention disclosure"
    document.core_properties.author = "OpenMycelium Project"
    document.core_properties.keywords = "heterogeneous GPU, Kubernetes, distributed training, MCCL, Mycelium"
    document.core_properties.comments = "Engineering preprint. Not legal advice or an official IEEE template."
    document.save(OUTPUT)
    build_pdf()
    print(OUTPUT)


if __name__ == "__main__":
    main()
